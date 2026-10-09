"""Tests for build_members_excel — sheet classification matches the exact
same member_matches_filter predicate GET /members?filter= uses, so the
report never drifts from what the app itself shows under those filters."""

from io import BytesIO

from openpyxl import load_workbook

from excel_report import COLUMNS, build_members_excel


def _member(mid, name, number, sub=None):
    return {"id": mid, "name": name, "phone": None, "email": None, "member_number": number, "current_subscription": sub}


def _sub(status="ACTIVE", end_date=None, due_amount=0, plan_name="Monthly"):
    return {"plan_name": plan_name, "status": status, "start_date": "2026-01-01", "end_date": end_date, "due_amount": due_amount}


def _load(xlsx_bytes):
    return load_workbook(BytesIO(xlsx_bytes))


class TestSheetStructure:
    def test_three_sheets_in_order(self):
        wb = _load(build_members_excel([]))
        assert wb.sheetnames == ["Active", "Expired", "Due"]

    def test_header_row_matches_columns(self):
        wb = _load(build_members_excel([]))
        header = next(wb["Active"].iter_rows(values_only=True))
        assert header == COLUMNS

    def test_empty_member_list_produces_header_only_sheets(self):
        wb = _load(build_members_excel([]))
        for name in wb.sheetnames:
            assert wb[name].max_row == 1


class TestClassification:
    def test_active_member_lands_in_active_sheet_only(self):
        members = [_member("m1", "Alice", 1, _sub(status="ACTIVE"))]
        wb = _load(build_members_excel(members))
        assert wb["Active"].max_row == 2
        assert wb["Expired"].max_row == 1
        assert wb["Due"].max_row == 1

    def test_expired_member_lands_in_expired_sheet(self):
        members = [_member("m1", "Bob", 2, _sub(status="EXPIRED"))]
        wb = _load(build_members_excel(members))
        assert wb["Active"].max_row == 1
        assert wb["Expired"].max_row == 2

    def test_member_with_due_amount_lands_in_due_regardless_of_status(self):
        """An expired member who still owes money belongs in Due too —
        same rule the app's own Due filter uses."""
        members = [_member("m1", "Carol", 3, _sub(status="EXPIRED", due_amount=500))]
        wb = _load(build_members_excel(members))
        assert wb["Expired"].max_row == 2
        assert wb["Due"].max_row == 2

    def test_member_can_appear_in_both_active_and_due(self):
        members = [_member("m1", "Dave", 4, _sub(status="ACTIVE", due_amount=200))]
        wb = _load(build_members_excel(members))
        assert wb["Active"].max_row == 2
        assert wb["Due"].max_row == 2

    def test_member_with_no_subscription_appears_nowhere(self):
        members = [_member("m1", "Eve", 5, None)]
        wb = _load(build_members_excel(members))
        for name in wb.sheetnames:
            assert wb[name].max_row == 1

    def test_frozen_member_appears_in_neither_active_nor_expired(self):
        members = [_member("m1", "Frank", 6, _sub(status="FROZEN"))]
        wb = _load(build_members_excel(members))
        assert wb["Active"].max_row == 1
        assert wb["Expired"].max_row == 1


class TestRowContent:
    def test_row_has_expected_values_in_order(self):
        members = [_member("m1", "Grace", 7, _sub(status="ACTIVE", end_date="2026-06-01", due_amount=0, plan_name="Annual"))]
        wb = _load(build_members_excel(members))
        row = list(wb["Active"].iter_rows(values_only=True))[1]
        # phone/email were written as "" (no phone on record) — openpyxl
        # round-trips an empty string as a blank cell (None), not "": a
        # storage-format detail, not a functional difference (both render
        # as an empty Excel cell either way).
        assert row == (7, "Grace", None, None, "Annual", "2026-01-01", "2026-06-01", 0, "ACTIVE")

    def test_expired_sheet_is_sorted_most_recently_lapsed_first(self):
        members = [
            _member("m1", "Old", 1, _sub(status="EXPIRED", end_date="2026-01-01")),
            _member("m2", "Recent", 2, _sub(status="EXPIRED", end_date="2026-09-01")),
        ]
        wb = _load(build_members_excel(members))
        names = [r[1] for r in list(wb["Expired"].iter_rows(values_only=True))[1:]]
        assert names == ["Recent", "Old"]
