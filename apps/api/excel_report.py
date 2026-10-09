"""Daily Active/Expired/Due Excel report — one workbook, one sheet per
section, classified with the exact same member_matches_filter() predicate
the Members page's own filters use, so "who's in the Expired sheet" never
drifts from "who shows up under the Expired filter in the app".
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from members import member_matches_filter, sort_members_for_filter

COLUMNS = ("Member ID", "Name", "Phone", "Email", "Plan", "Start Date", "End Date", "Due Amount", "Status")

HEADER_FILL = PatternFill(start_color="FFD60A", end_color="FFD60A", fill_type="solid")  # brand yellow
HEADER_FONT = Font(bold=True)


def _row_for(member: dict) -> tuple:
    sub = member.get("current_subscription")
    return (
        member.get("member_number"),
        member["name"],
        member.get("phone") or "",
        member.get("email") or "",
        sub["plan_name"] if sub else "",
        sub["start_date"] if sub else "",
        sub["end_date"] if sub and sub.get("end_date") else "",
        sub["due_amount"] if sub else 0,
        sub["status"] if sub else "No plan",
    )


def _write_sheet(ws: Worksheet, title: str, members: list[dict]) -> None:
    ws.title = title
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for m in members:
        ws.append(_row_for(m))
    for col_cells in ws.columns:
        width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 40)


def build_members_excel(members: list[dict]) -> bytes:
    """`members` — the same enriched shape GET /members returns (each with
    an optional `current_subscription` dict already attached). Classifies
    into three sheets using the live member_matches_filter predicate:
      - Active: current_subscription.status == ACTIVE
      - Expired: current_subscription.status == EXPIRED (sorted most
        recently lapsed first, matching the Expired filter's own order)
      - Due: due_amount > 0, regardless of status (an expired member who
        still owes money belongs here too — same rule the Due filter uses)
    A member can legitimately appear in both Expired and Due.
    """
    latest_by_member = {m["id"]: m.get("current_subscription") for m in members}

    active = [m for m in members if member_matches_filter(m.get("current_subscription"), "active")]
    expired = [m for m in members if member_matches_filter(m.get("current_subscription"), "expired")]
    expired = sort_members_for_filter(expired, latest_by_member, "expired")
    due = [m for m in members if member_matches_filter(m.get("current_subscription"), "due")]

    wb = Workbook()
    _write_sheet(wb.active, "Active", active)
    _write_sheet(wb.create_sheet(), "Expired", expired)
    _write_sheet(wb.create_sheet(), "Due", due)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
